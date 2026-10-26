"""Read only the filing allowlist. Evaluation labels are never loaded here."""

import csv
import hashlib
import json
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from pydantic import ValidationError

from backend.app.schemas import Filing, FilingMetadata, RateRow, SourceRow

MAX_ROWS = 5000
MAX_FILE_BYTES = 2_000_000
MAX_XLSX_EXPANDED_BYTES = 20_000_000
REQUIRED_COLUMNS = {
    "segment_id", "current_base_rate", "proposed_base_rate",
    "current_rating_factor", "proposed_rating_factor",
}


class InputValidationError(ValueError):
    def __init__(self, issues: list[dict]):
        self.issues = issues
        super().__init__("; ".join(f"{i['location']}: {i['message']}" for i in issues))


def fail(location: str, message: str):
    raise InputValidationError([{"location": location, "message": message}])


def validation_issues(error: ValidationError, location: str) -> list[dict]:
    # Do not include raw input values in errors.
    return [
        {"location": f"{location}.{'.'.join(map(str, item['loc']))}", "message": item["msg"]}
        for item in error.errors(include_input=False, include_url=False)
    ]


def check_header(header: list, filename: str):
    if not header or any(not isinstance(col, str) or not col for col in header):
        fail(filename, "A nonempty string header is required for every column")
    if len(header) != len(set(header)):
        fail(filename, "Duplicate column headers")
    missing = REQUIRED_COLUMNS - set(header)
    extra = set(header) - set(RateRow.model_fields)
    if missing:
        fail(filename, f"Missing required columns: {', '.join(sorted(missing))}")
    if extra:
        fail(filename, f"Unknown columns: {', '.join(sorted(extra))}")


def validate_records(records: list[tuple[int, dict]], filename="rates.csv", sheet=None):
    if not records:
        fail(filename, "At least one data row is required")
    if len(records) > MAX_ROWS:
        fail(filename, f"Maximum {MAX_ROWS} data rows exceeded")
    rows, sources, issues = [], [], []
    seen = set()
    for row_number, record in records:
        normalized = {key: (None if value is None or value == "" else value) for key, value in record.items()}
        # Empty descriptive fields are allowed; numeric blanks remain missing.
        for key in ("territory", "driver_class", "vehicle_group"):
            if key in normalized and normalized[key] is None:
                normalized[key] = ""
        try:
            row = RateRow.model_validate(normalized)
        except ValidationError as error:
            issues.extend(validation_issues(error, f"{filename}:{row_number}"))
            continue
        if row.segment_id in seen:
            issues.append({"location": f"{filename}:{row_number}.segment_id", "message": "Duplicate segment ID"})
        seen.add(row.segment_id)
        rows.append(row)
        sources.append(SourceRow(evidence_id=f"row:{row.segment_id}", file=filename, row_number=row_number, sheet=sheet))
    if issues:
        raise InputValidationError(issues)
    return rows, sources


def load_rates(path: Path):
    path = Path(path)
    if path.stat().st_size > MAX_FILE_BYTES:
        fail(path.name, "File exceeds 2 MB limit")
    records = []
    if path.suffix.lower() == ".csv":
        try:
            with path.open(encoding="utf-8-sig", newline="") as handle:
                reader = csv.reader(handle, strict=True)
                header = next(reader, [])
                check_header(header, path.name)
                while True:
                    row_number = reader.line_num + 1
                    try:
                        values = next(reader)
                    except StopIteration:
                        break
                    if not values:
                        continue
                    if len(values) != len(header):
                        fail(f"{path.name}:{row_number}", "Row width does not match header")
                    records.append((row_number, dict(zip(header, values))))
                    if len(records) > MAX_ROWS:
                        fail(path.name, f"Maximum {MAX_ROWS} data rows exceeded")
        except (csv.Error, UnicodeError):
            fail(path.name, "Invalid UTF-8 CSV")
        return validate_records(records, path.name)
    if path.suffix.lower() != ".xlsx":
        fail(path.name, "Only CSV and values-only XLSX files are supported")
    workbook = None
    try:
        with ZipFile(path) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_XLSX_EXPANDED_BYTES:
                fail(path.name, "Expanded XLSX exceeds 20 MB limit")
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        if "rates" not in workbook.sheetnames:
            fail(path.name, "XLSX requires a sheet named rates")
        sheet = workbook["rates"]
        if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > 32:
            fail(path.name, "XLSX sheet dimensions exceed limits")
        iterator = sheet.iter_rows()
        header = [cell.value for cell in next(iterator, [])]
        check_header(header, path.name)
        for row_number, cells in enumerate(iterator, start=2):
            if any(cell.data_type == "f" for cell in cells):
                fail(f"{path.name}:{row_number}", "Formulas are unsupported; export values only")
            if all(cell.value is None for cell in cells):
                continue
            records.append((row_number, dict(zip(header, [cell.value for cell in cells]))))
            if len(records) > MAX_ROWS:
                fail(path.name, f"Maximum {MAX_ROWS} data rows exceeded")
        return validate_records(records, path.name, "rates")
    except InputValidationError:
        raise
    except (BadZipFile, KeyError, OSError, ParseError, ValueError):
        fail(path.name, "Invalid XLSX workbook")
    finally:
        if workbook:
            workbook.close()


def load_filing(directory: Path, rates_filename: str = "rates.csv") -> Filing:
    """Explicit allowlist prevents answer keys or unrelated files entering context."""
    directory = Path(directory)
    if rates_filename not in {"rates.csv", "rates.xlsx"}:
        fail("rates_filename", "Choose rates.csv or rates.xlsx")
    paths = {name: directory / name for name in (rates_filename, "filing.json", "narrative.md")}
    for name, path in paths.items():
        if not path.is_file():
            fail(name, "Required filing input not found")
        if path.stat().st_size > MAX_FILE_BYTES:
            fail(name, "File exceeds 2 MB limit")
    rows, sources = load_rates(paths[rates_filename])
    try:
        metadata = FilingMetadata.model_validate_json(paths["filing.json"].read_text(encoding="utf-8"))
        narrative = paths["narrative.md"].read_text(encoding="utf-8")
        return Filing(
            metadata=metadata, rows=rows, source_rows=sources, narrative=narrative,
            file_sha256={name: hashlib.sha256(path.read_bytes()).hexdigest() for name, path in paths.items()},
        )
    except ValidationError as error:
        raise InputValidationError(validation_issues(error, directory.name)) from None
    except (UnicodeError, json.JSONDecodeError):
        fail(directory.name, "Invalid UTF-8 or JSON filing input")
